from __future__ import annotations

import csv
import io
import json
import os
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal

import pyodbc
from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from analysis import router as analysis_router
from db import get_connection

BASE_DIR = Path(__file__).resolve().parent
FIELD_MAP_PATH = BASE_DIR / "field_map.json"
QUERY_RULES_PATH = BASE_DIR / "query_rules.json"

app = FastAPI(title="庫存分析與明細查詢", version="14.9.0")
app.include_router(analysis_router)


@app.middleware("http")
async def disable_browser_cache(request: Request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")

# 庫存明細查詢沿用 v12，不改動資料來源、篩選、分頁與匯出邏輯。
DETAIL_VIEW_MAP: dict[str, str] = {
    # ALL 直接查目前正式共用 BASE VIEW；各指定類型仍查個別 Type VIEW。
    "ALL": "dbo.vw_InventoryInquiry_Base",
    "1": "dbo.vw_InventoryInquiry_Type189",
    "2": "dbo.vw_InventoryInquiry_Type2",
    "3": "dbo.vw_InventoryInquiry_Type3",
    "4": "dbo.vw_InventoryInquiry_Type4",
    "8": "dbo.vw_InventoryInquiry_Type189",
    "9": "dbo.vw_InventoryInquiry_Type189",
    "D": "dbo.vw_InventoryInquiry_TypeD",
    "E": "dbo.vw_InventoryInquiry_TypeE",
    "F": "dbo.vw_InventoryInquiry_TypeF",
}

SHARED_TYPE_CODES = {"1", "8", "9"}

TYPE_LABELS: dict[str, str] = {
    "ALL": "全部",
    "1": "布",
    "2": "胚布",
    "3": "紗",
    "4": "盤頭紗",
    "8": "成衣",
    "9": "領片",
    "D": "代工",
    "E": "膜",
    "F": "軍工／其他",
}

GREIGE_TYPE_LABELS: dict[str, str] = {
    "1": "針織",
    "2": "平織",
    "3": "領片胚布",
    "4": "不織布",
}

HIDDEN_RESULT_COLUMNS = {
    "庫存識別碼",
    "載入批次ID",
    "快照時間",
    "計算基準日",
    "庫存鍵",
    "InventoryKey",
    "SKU",
}

DETAIL_FIELD_CANDIDATES: dict[str, list[str]] = {
    "operating_center": ["營運中心", "OperatingCenter"],
    "type_discriminator": ["類型"],
    "item_no": ["料號"],
    "warehouse": ["倉庫"],
    "color_no": ["興采色號"],
    "zero_project": ["零元專案否"],
    "project_302": ["302專案否", "302專案", "Project302Flag"],
    "project_batch_no": ["批號(訂單/請購單)", "imgs04"],
    "salesperson": ["業務姓名", "業務"],
    "grade": ["等級"],
    "reason_code": ["理由碼"],
    "customer": ["Final Buyer名稱", "Final Buyer ID"],
    "business_department": ["業務部門簡稱", "業務部門"],
    "order_type": ["訂單類別", "訂單類別代碼"],
    "greige_type": ["胚布類型", "ta_ima20"],
    # 保留 API 欄位名稱 inbound_date，實際對應 4GL 使用的庫存異動日期。
    "inbound_date": ["異動日期"],
}

# 可輸入搜尋下拉選單允許查詢的欄位。
# 前端只能傳入此白名單中的 key，實際 SQL 欄位仍由 VIEW 欄位偵測結果決定。
DETAIL_OPTION_FIELDS = {
    "item_no",
    "warehouse",
    "color_no",
    "zero_project",
    "salesperson",
    "business_department",
    "grade",
    "reason_code",
    "customer",
    "order_type",
    "greige_type",
}


class DetailSearchRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    inventory_type: Literal["ALL", "1", "2", "3", "4", "8", "9", "D", "E", "F"] = Field(alias="type")
    operating_centers: list[str] = Field(default_factory=list)
    item_nos: list[str] = Field(default_factory=list)
    warehouses: list[str] = Field(default_factory=list)
    color_nos: list[str] = Field(default_factory=list)
    zero_projects: list[str] = Field(default_factory=list)
    project_302_flags: list[str] = Field(default_factory=list)
    salespersons: list[str] = Field(default_factory=list)
    business_departments: list[str] = Field(default_factory=list)
    grades: list[str] = Field(default_factory=list)
    reason_codes: list[str] = Field(default_factory=list)
    customers: list[str] = Field(default_factory=list)
    order_types: list[str] = Field(default_factory=list)
    greige_types: list[str] = Field(default_factory=list)
    inbound_date_start: date | None = None
    inbound_date_end: date | None = None
    page: int = 1
    page_size: int = 100
    sort_field: str | None = None
    sort_direction: Literal["ASC", "DESC"] = "DESC"


def qident(name: str) -> str:
    return "[" + name.replace("]", "]]" ) + "]"


def split_view_name(view_name: str) -> tuple[str, str]:
    parts = view_name.split(".", maxsplit=1)
    if len(parts) != 2:
        raise RuntimeError(f"VIEW 名稱格式錯誤：{view_name}")
    return parts[0], parts[1]


def load_field_overrides() -> dict[str, dict[str, str]]:
    if not FIELD_MAP_PATH.exists():
        return {}
    with FIELD_MAP_PATH.open("r", encoding="utf-8") as fp:
        data = json.load(fp)
    if not isinstance(data, dict):
        raise RuntimeError("field_map.json 格式錯誤。")
    return data



def load_query_rules() -> dict[str, Any]:
    """
    讀取查詢規則設定檔。

    不使用快取，讓 query_rules.json 修改後可直接於下一次查詢生效，
    不需要為了調整排除清單而修改 Python 程式。
    """
    if not QUERY_RULES_PATH.exists():
        return {}

    try:
        with QUERY_RULES_PATH.open("r", encoding="utf-8") as fp:
            data = json.load(fp)
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(
            status_code=500,
            detail=f"query_rules.json 讀取失敗：{exc}",
        ) from exc

    if not isinstance(data, dict):
        raise HTTPException(
            status_code=500,
            detail="query_rules.json 格式錯誤：根節點必須是 JSON object。",
        )
    return data


def get_detail_excluded_values(field: str) -> list[str]:
    """
    取得庫存明細查詢固定排除值。

    設定位置：
    detail_query.exclude.<field>
    """
    rules = load_query_rules()
    detail_query = rules.get("detail_query", {})
    exclude = detail_query.get("exclude", {})

    if not isinstance(detail_query, dict) or not isinstance(exclude, dict):
        return []

    rule = exclude.get(field, {})
    if not isinstance(rule, dict):
        return []

    if rule.get("enabled", True) is False:
        return []

    values = rule.get("values", [])
    if not isinstance(values, list):
        return []

    result: list[str] = []
    for value in values:
        clean = str(value).strip()
        if clean:
            result.append(clean)

    return result



def get_project_302_force_y_batch_numbers() -> list[str]:
    """
    取得 query_rules.json 中，沒有固定公式但必須視為 302 專案的批號。
    """
    rules = load_query_rules()
    detail_query = rules.get("detail_query", {})
    overrides = detail_query.get("project_302_overrides", {})

    if not isinstance(detail_query, dict) or not isinstance(overrides, dict):
        return []

    if overrides.get("enabled", True) is False:
        return []

    values = overrides.get("force_y_batch_numbers", [])
    if not isinstance(values, list):
        return []

    return [
        str(value).strip()
        for value in values
        if str(value).strip()
    ]


def sql_nvarchar_literal(value: str) -> str:
    return "N'" + str(value).replace("'", "''") + "'"


def project_302_effective_expression(
    mapping: dict[str, str | None],
) -> str | None:
    """
    有效 302 專案判斷：
    1. query_rules.json 特例批號 -> Y
    2. 其他 -> 使用資料來源原本的 302 專案值
    """
    project_col = mapping.get("project_302")
    if not project_col:
        return None

    raw_expression = (
        f"UPPER(LTRIM(RTRIM(CONVERT(nvarchar(20), "
        f"{qident(project_col)}))))"
    )

    batch_numbers = get_project_302_force_y_batch_numbers()
    batch_col = mapping.get("project_batch_no")

    if not batch_numbers or not batch_col:
        return raw_expression

    batch_expression = (
        f"UPPER(LTRIM(RTRIM(CONVERT(nvarchar(200), "
        f"{qident(batch_col)}))))"
    )
    literals = ", ".join(
        sql_nvarchar_literal(value.upper())
        for value in batch_numbers
    )

    return (
        "CASE "
        f"WHEN {batch_expression} IN ({literals}) THEN N'Y' "
        f"ELSE {raw_expression} "
        "END"
    )


def add_project_302_filter(
    where: list[str],
    params: list[Any],
    mapping: dict[str, str | None],
    values: list[str],
) -> None:
    """
    使用「含 Config 特例」後的有效 302 專案值篩選。
    """
    clean = [
        str(value).strip().upper()
        for value in values
        if str(value).strip()
    ]
    if not clean:
        return

    expression = project_302_effective_expression(mapping)
    if not expression:
        raise HTTPException(
            status_code=400,
            detail="目前資料來源沒有「302專案」欄位。",
        )

    placeholders = ", ".join("?" for _ in clean)
    where.append(f"{expression} IN ({placeholders})")
    params.extend(clean)


def build_detail_select_list(
    visible_columns: list[str],
    mapping: dict[str, str | None],
) -> str:
    """
    查詢結果中的 302 專案欄位，也套用 Config 特例，
    使畫面 / Excel / CSV 顯示一致。
    """
    project_col = mapping.get("project_302")
    project_expression = project_302_effective_expression(mapping)

    parts: list[str] = []
    for column in visible_columns:
        if (
            project_col
            and project_expression
            and column == project_col
        ):
            parts.append(
                f"{project_expression} AS {qident(column)}"
            )
        else:
            parts.append(qident(column))

    return ", ".join(parts)


def add_config_exclusion_filters(
    where: list[str],
    params: list[Any],
    mapping: dict[str, str | None],
) -> None:
    """
    套用 query_rules.json 的庫存明細固定排除規則。

    目前先支援 warehouse；未來可在此集中擴充其他欄位。
    """
    warehouse_values = get_detail_excluded_values("warehouse")
    if warehouse_values:
        warehouse_col = mapping.get("warehouse")
        if warehouse_col:
            placeholders = ", ".join("?" for _ in warehouse_values)
            expression = (
                f"UPPER(LTRIM(RTRIM(CONVERT(nvarchar(4000), "
                f"{qident(warehouse_col)}))))"
            )
            where.append(
                f"{expression} NOT IN ({placeholders})"
            )
            params.extend(value.upper() for value in warehouse_values)


def normalize(value: str) -> str:
    return re.sub(r"[\s_\-（）()\[\]【】/\\]+", "", value).casefold()


def detect_column(columns: list[str], candidates: list[str]) -> str | None:
    normalized_columns = {normalize(col): col for col in columns}
    for candidate in candidates:
        exact = normalized_columns.get(normalize(candidate))
        if exact:
            return exact

    for candidate in candidates:
        normalized_candidate = normalize(candidate)
        for column in columns:
            normalized_column = normalize(column)
            if normalized_candidate and (
                normalized_candidate in normalized_column
                or normalized_column in normalized_candidate
            ):
                return column
    return None


@lru_cache(maxsize=64)
def get_view_columns(view_name: str) -> tuple[dict[str, Any], ...]:
    schema, object_name = split_view_name(view_name)
    sql = """
    SELECT
        c.column_id,
        c.name AS column_name,
        t.name AS data_type,
        c.max_length,
        c.precision,
        c.scale,
        c.is_nullable,
        o.type_desc AS object_type
    FROM sys.objects AS o
    INNER JOIN sys.schemas AS s
        ON s.schema_id = o.schema_id
    INNER JOIN sys.columns AS c
        ON c.object_id = o.object_id
    INNER JOIN sys.types AS t
        ON t.user_type_id = c.user_type_id
    WHERE s.name = ?
      AND o.name = ?
      AND o.type IN (N'U', N'V')
    ORDER BY c.column_id;
    """
    with get_connection() as conn:
        rows = conn.cursor().execute(sql, schema, object_name).fetchall()

    if not rows:
        raise HTTPException(status_code=404, detail=f"找不到 InventoryInquiry 資料來源：{view_name}")

    return tuple(
        {
            "ordinal": int(row.column_id),
            "name": str(row.column_name),
            "data_type": str(row.data_type),
            "max_length": int(row.max_length),
            "precision": int(row.precision),
            "scale": int(row.scale),
            "nullable": bool(row.is_nullable),
            "object_type": str(row.object_type),
        }
        for row in rows
    )


def visible_metadata(view_name: str) -> list[dict[str, Any]]:
    metadata = [
        dict(item)
        for item in get_view_columns(view_name)
        if item["name"] not in HIDDEN_RESULT_COLUMNS
    ]

    # 營運中心固定顯示為查詢結果第一欄。
    # 新版 VIEW 若已提供 [營運中心] / [OperatingCenter] 就直接使用；
    # 舊版 VIEW 尚未提供時，由 detail_source_sql() 從
    # InventoryInquirySnapshot_Work 補出 [營運中心]。
    operating_center_column = detect_column(
        [item["name"] for item in metadata],
        ["營運中心", "OperatingCenter"],
    )
    if operating_center_column:
        center_item = next(
            item for item in metadata
            if item["name"] == operating_center_column
        )
    else:
        center_item = {
            "ordinal": 0,
            "name": "營運中心",
            "data_type": "nvarchar",
            "max_length": 40,
            "precision": 0,
            "scale": 0,
            "nullable": True,
            "object_type": "SYNTHETIC",
        }

    center_item["label"] = "營運中心"
    metadata = [
        center_item,
        *[
            item for item in metadata
            if item["name"] != operating_center_column
            and item["name"] != "營運中心"
        ],
    ]

    # 明細欄位依使用者指定位置重新排列；未指定欄位維持 VIEW 原順序。
    # 不同 Type View 若缺少個別欄位，會直接略過，不影響其他欄位。
    item_after_activity_date = ["料號"]
    fields_after_location_name = [
        "訂單類別代碼",
        "訂單類別",
        "Final Buyer ID",
        "Final Buyer名稱",
        "中文色名",
        "英文色名",
        "興采色號",
    ]
    ending_fields = [
        "規格",
        "品名",
        "零元專案否",
        "Tranzend上架品名",
    ]

    moved_names = {
        *item_after_activity_date,
        *fields_after_location_name,
        *ending_fields,
    }
    item_by_name = {item["name"]: item for item in metadata}
    ordered = [item for item in metadata if item["name"] not in moved_names]

    def insert_after(anchor_name: str, field_names: list[str]) -> None:
        items = [item_by_name[name] for name in field_names if name in item_by_name]
        if not items:
            return

        anchor_index = next(
            (index for index, item in enumerate(ordered) if item["name"] == anchor_name),
            None,
        )
        if anchor_index is None:
            ordered.extend(items)
        else:
            ordered[anchor_index + 1:anchor_index + 1] = items

    insert_after("異動日期", item_after_activity_date)
    insert_after("儲位名稱", fields_after_location_name)
    ordered.extend(item_by_name[name] for name in ending_fields if name in item_by_name)
    return ordered


def detail_source_sql(view_name: str) -> tuple[str, list[str]]:
    """
    回傳明細查詢實際 FROM 來源與可用欄位。

    個別 Type VIEW / 統一 VIEW 的欄位並不完全一致。
    若目前資料來源缺少下列欄位，利用 InventoryInquirySnapshot_Work.InventoryKey
    補出查詢所需的固定欄位，不必逐支修改既有 VIEW：

    - [營運中心]  <- Work.OperatingCenter
    - [302專案否] <- Work.Project302Flag

    這樣「類型=全部」使用 dbo.vw_InventoryInquiry_Base 時，
    302 專案 Y/N 條件仍可正常查詢。
    """

    view_columns = [item["name"] for item in get_view_columns(view_name)]

    operating_center_column = detect_column(
        view_columns,
        ["營運中心", "OperatingCenter"],
    )
    project_302_column = detect_column(
        view_columns,
        ["302專案否", "302專案", "Project302Flag"],
    )

    need_operating_center = operating_center_column is None
    need_project_302 = project_302_column is None

    # VIEW 本身欄位完整時，不增加 JOIN。
    if not need_operating_center and not need_project_302:
        return view_name, view_columns

    if "庫存識別碼" not in view_columns:
        missing_labels = []
        if need_operating_center:
            missing_labels.append("營運中心")
        if need_project_302:
            missing_labels.append("302專案否")
        raise HTTPException(
            status_code=500,
            detail=(
                f"{view_name} 缺少「庫存識別碼」，"
                f"無法從 InventoryInquirySnapshot_Work 補出："
                + "、".join(missing_labels)
            ),
        )

    synthetic_selects: list[str] = []
    synthetic_columns: list[str] = []

    if need_operating_center:
        synthetic_selects.append(
            "LTRIM(RTRIM(CONVERT(nvarchar(20), w.OperatingCenter))) "
            "AS [營運中心]"
        )
        synthetic_columns.append("營運中心")

    if need_project_302:
        synthetic_selects.append(
            "LTRIM(RTRIM(CONVERT(nvarchar(10), w.Project302Flag))) "
            "AS [302專案否]"
        )
        synthetic_columns.append("302專案否")

    source_sql = (
        "(SELECT v.*, "
        + ", ".join(synthetic_selects)
        + f" FROM {view_name} AS v "
        "INNER JOIN dbo.InventoryInquirySnapshot_Work AS w "
        "ON w.InventoryKey = "
        "TRY_CONVERT(varbinary(32), v.[庫存識別碼], 2)) AS detail_source"
    )

    return source_sql, [*view_columns, *synthetic_columns]


def resolve_detail_mapping(view_name: str, columns: list[str]) -> dict[str, str | None]:
    overrides = load_field_overrides().get(view_name, {})
    mapping: dict[str, str | None] = {}

    for key, candidates in DETAIL_FIELD_CANDIDATES.items():
        override = str(overrides.get(key, "") or "").strip()
        if override:
            if override not in columns:
                raise HTTPException(
                    status_code=500,
                    detail=f"{view_name} 的 field_map.json 欄位不存在：{key}={override}",
                )
            mapping[key] = override
        else:
            mapping[key] = detect_column(columns, candidates)
    return mapping


def require_field(mapping: dict[str, str | None], key: str, label: str) -> str:
    column = mapping.get(key)
    if not column:
        raise HTTPException(status_code=400, detail=f"資料來源無法辨識「{label}」欄位。")
    return column


def serialize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return value



def excel_cell_value(value: Any) -> Any:
    """
    Excel 匯出專用型別轉換。

    重點：
    - Decimal / int / float 保持數值型別，Excel 才能加總、排序、公式運算。
    - date / datetime 保持日期型別。
    - 只有 bytes 轉成十六進位文字。
    - 原本就是字串的料號、單號、項次代碼等仍保持文字，
      不會因為「看起來像數字」就被強制轉成數字。
    """
    if value is None:
        return None
    if isinstance(value, bytes):
        return value.hex()
    return value


def result_from_cursor(cursor: pyodbc.Cursor) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    names = [desc[0] for desc in cursor.description]
    rows = [
        {name: serialize_value(value) for name, value in zip(names, row)}
        for row in cursor.fetchall()
    ]
    columns = [{"ordinal": index + 1, "name": name} for index, name in enumerate(names)]
    return columns, rows


def translate_detail_classification_values(
    rows: list[dict[str, Any]],
    mapping: dict[str, str | None],
) -> None:
    """將明細結果中的分類代碼轉成中文，不改動資料庫 VIEW 原值。"""

    type_column = mapping.get("type_discriminator")
    greige_type_column = mapping.get("greige_type")

    for row in rows:
        if type_column and type_column in row:
            raw_type = row.get(type_column)
            if raw_type is not None:
                type_code = str(raw_type).strip().upper()
                row[type_column] = TYPE_LABELS.get(type_code, raw_type)

        if greige_type_column and greige_type_column in row:
            raw_greige_type = row.get(greige_type_column)
            if raw_greige_type is not None:
                greige_type_code = str(raw_greige_type).strip()
                row[greige_type_column] = GREIGE_TYPE_LABELS.get(
                    greige_type_code,
                    raw_greige_type,
                )


def add_multi_filter(
    where: list[str],
    params: list[Any],
    column: str | None,
    values: list[str],
    label: str,
    *,
    binary_compare: bool = False,
) -> None:
    clean = [str(v).strip() for v in values if str(v).strip()]
    if not clean:
        return
    if not column:
        raise HTTPException(status_code=400, detail=f"目前資料來源沒有「{label}」欄位。")

    placeholders = ", ".join("?" for _ in clean)
    expression = (
        f"LTRIM(RTRIM(CONVERT(nvarchar(4000), {qident(column)})))"
    )

    if binary_compare:
        expression += " COLLATE Latin1_General_100_BIN2"

    where.append(f"{expression} IN ({placeholders})")
    params.extend(clean)


def validate_paging(req: DetailSearchRequest) -> None:
    max_page_size = int(os.getenv("MAX_PAGE_SIZE", "500"))
    if req.page < 1:
        raise HTTPException(status_code=400, detail="page 必須大於等於 1。")
    if req.page_size < 1 or req.page_size > max_page_size:
        raise HTTPException(
            status_code=400,
            detail=f"page_size 必須介於 1 與 {max_page_size}。",
        )


def build_detail_filter_sql(
    req: DetailSearchRequest,
    mapping: dict[str, str | None],
) -> tuple[str, list[Any]]:
    where: list[str] = ["1 = 1"]
    params: list[Any] = []

    # 先套用固定排除規則，再套用使用者查詢條件。
    add_config_exclusion_filters(where, params, mapping)

    add_multi_filter(
        where,
        params,
        mapping.get("operating_center"),
        req.operating_centers,
        "營運中心",
    )

    if req.inventory_type in SHARED_TYPE_CODES:
        type_col = require_field(mapping, "type_discriminator", "類型")
        where.append(f"CONVERT(nvarchar(50), {qident(type_col)}) = ?")
        params.append(req.inventory_type)

    add_multi_filter(where, params, mapping.get("item_no"), req.item_nos, "料號")
    add_multi_filter(where, params, mapping.get("warehouse"), req.warehouses, "倉庫")
    add_multi_filter(
        where,
        params,
        mapping.get("color_no"),
        req.color_nos,
        "興采色號",
        binary_compare=True,
    )
    add_multi_filter(
        where,
        params,
        mapping.get("zero_project"),
        req.zero_projects,
        "零元專案否",
    )
    add_project_302_filter(
        where,
        params,
        mapping,
        req.project_302_flags,
    )
    add_multi_filter(where, params, mapping.get("salesperson"), req.salespersons, "業務")
    add_multi_filter(
        where,
        params,
        mapping.get("business_department"),
        req.business_departments,
        "業務部門",
    )
    add_multi_filter(where, params, mapping.get("grade"), req.grades, "等級")
    add_multi_filter(
        where,
        params,
        mapping.get("reason_code"),
        req.reason_codes,
        "理由碼",
    )
    add_multi_filter(where, params, mapping.get("customer"), req.customers, "客戶")
    add_multi_filter(
        where,
        params,
        mapping.get("order_type"),
        req.order_types,
        "訂單類別",
    )

    if req.inventory_type == "2":
        add_multi_filter(
            where,
            params,
            mapping.get("greige_type"),
            req.greige_types,
            "胚布類型",
        )

    if req.inbound_date_start or req.inbound_date_end:
        inbound_col = require_field(mapping, "inbound_date", "異動日期")
        if req.inbound_date_start:
            where.append(f"{qident(inbound_col)} >= CAST(? AS date)")
            params.append(req.inbound_date_start.isoformat())
        if req.inbound_date_end:
            where.append(f"{qident(inbound_col)} < DATEADD(day, 1, CAST(? AS date))")
            params.append(req.inbound_date_end.isoformat())

    return " AND ".join(where), params


def fetch_detail_rows(req: DetailSearchRequest, *, export: bool = False) -> dict[str, Any]:
    validate_paging(req)

    view_name = DETAIL_VIEW_MAP[req.inventory_type]
    source_sql, all_columns = detail_source_sql(view_name)
    metadata = visible_metadata(view_name)
    visible_columns = [item["name"] for item in metadata]
    mapping = resolve_detail_mapping(view_name, all_columns)
    where_sql, params = build_detail_filter_sql(req, mapping)

    if req.sort_field and req.sort_field not in visible_columns:
        raise HTTPException(status_code=400, detail="排序欄位不屬於目前顯示欄位。")

    default_sort = mapping.get("inbound_date")
    if default_sort not in visible_columns:
        default_sort = visible_columns[0]
    sort_field = req.sort_field or default_sort

    select_list = build_detail_select_list(visible_columns, mapping)

    if export:
        max_export_rows = int(os.getenv("MAX_EXPORT_ROWS", "100000"))
        sql = (
            f"SELECT TOP ({max_export_rows}) {select_list} "
            f"FROM {source_sql} "
            f"WHERE {where_sql} "
            f"ORDER BY {qident(sort_field)} {req.sort_direction};"
        )
        query_params = params
    else:
        offset = (req.page - 1) * req.page_size
        sql = (
            f"SELECT {select_list} "
            f"FROM {source_sql} "
            f"WHERE {where_sql} "
            f"ORDER BY {qident(sort_field)} {req.sort_direction} "
            f"OFFSET ? ROWS FETCH NEXT ? ROWS ONLY;"
        )
        query_params = [*params, offset, req.page_size]

    count_sql = f"SELECT COUNT_BIG(*) FROM {source_sql} WHERE {where_sql};"

    try:
        with get_connection() as conn:
            total_rows = int(conn.cursor().execute(count_sql, *params).fetchval())
            cursor = conn.cursor().execute(sql, *query_params)
            _, rows = result_from_cursor(cursor)
            translate_detail_classification_values(rows, mapping)
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"SQL Server 查詢失敗：{exc}") from exc

    return {
        "type": req.inventory_type,
        "view_name": view_name,
        "total_rows": total_rows,
        "page": req.page,
        "page_size": req.page_size,
        "columns": metadata,
        "mapping": mapping,
        "rows": rows,
    }


@app.get("/", response_class=HTMLResponse)
def index(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "type_options": [
                {"code": code, "label": TYPE_LABELS[code]}
                for code in DETAIL_VIEW_MAP.keys()
            ],
            "greige_type_options": [
                {"code": code, "label": label}
                for code, label in GREIGE_TYPE_LABELS.items()
            ],
        },
    )


@app.get("/api/health")
def health() -> dict[str, Any]:
    try:
        with get_connection() as conn:
            value = conn.cursor().execute("SELECT 1;").fetchval()
        return {
            "ok": value == 1,
            "app_version": app.version,
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@app.get("/api/detail/schema")
def detail_schema(type: Literal["ALL", "1", "2", "3", "4", "8", "9", "D", "E", "F"] = Query(...)) -> dict[str, Any]:
    view_name = DETAIL_VIEW_MAP[type]
    _, all_columns = detail_source_sql(view_name)
    metadata = visible_metadata(view_name)
    mapping = resolve_detail_mapping(view_name, all_columns)
    return {
        "type": type,
        "view_name": view_name,
        "columns": metadata,
        "mapping": mapping,
        "greige_type_enabled": type == "2",
        "hidden_columns": sorted(HIDDEN_RESULT_COLUMNS),
        "type_label": TYPE_LABELS[type],
        "greige_type_labels": GREIGE_TYPE_LABELS,
    }


def escape_like_prefix(value: str) -> str:
    """跳脫 SQL Server LIKE 萬用字元，供前綴查詢使用。"""

    return (
        value.replace("~", "~~")
        .replace("%", "~%")
        .replace("_", "~_")
        .replace("[", "~[")
    )


@app.get("/api/detail/option-search")
def detail_option_search(
    type: Literal["ALL", "1", "2", "3", "4", "8", "9", "D", "E", "F"] = Query(...),
    field: str = Query(..., min_length=1, max_length=50),
    keyword: str = Query(default="", max_length=200),
    operating_center: str = Query(default="", max_length=20),
    limit: int = Query(default=100, ge=1, le=200),
) -> dict[str, Any]:
    """
    提供明細查詢的可輸入下拉選項。

    搜尋規則：
    - 空白：顯示排序後前 N 筆。
    - 輸入 A：顯示 A 開頭。
    - 再輸入 AB：顯示 AB 開頭。
    - 僅允許白名單欄位，避免使用者直接指定 SQL 欄位。
    """

    if field not in DETAIL_OPTION_FIELDS:
        raise HTTPException(status_code=400, detail=f"不支援的選項欄位：{field}")
    if field == "greige_type" and type != "2":
        return {
            "type": type,
            "field": field,
            "keyword": keyword,
            "options": [],
        }

    view_name = DETAIL_VIEW_MAP[type]
    source_sql, all_columns = detail_source_sql(view_name)
    mapping = resolve_detail_mapping(view_name, all_columns)
    column = mapping.get(field)
    if not column:
        return {
            "type": type,
            "field": field,
            "keyword": keyword,
            "options": [],
        }

    clean_keyword = keyword.strip()
    clean_operating_center = operating_center.strip().upper()
    if clean_operating_center not in {"", "SNX", "GFUN", "M302"}:
        raise HTTPException(status_code=400, detail="不支援的營運中心。")

    option_expression = (
        f"LTRIM(RTRIM(CONVERT(nvarchar(4000), {qident(column)})))"
    )
    if field == "color_no":
        option_expression += " COLLATE Latin1_General_100_BIN2"

    where = [
        f"{qident(column)} IS NOT NULL",
        f"{option_expression} <> N''",
    ]
    params: list[Any] = []

    # 下拉選項亦依固定排除規則產生，避免 D35 專屬值出現在選項中。
    add_config_exclusion_filters(where, params, mapping)

    if clean_operating_center:
        operating_center_col = require_field(
            mapping,
            "operating_center",
            "營運中心",
        )
        where.append(
            f"LTRIM(RTRIM(CONVERT(nvarchar(50), {qident(operating_center_col)}))) = ?"
        )
        params.append(clean_operating_center)

    if type in SHARED_TYPE_CODES:
        type_col = require_field(mapping, "type_discriminator", "類型")
        where.append(f"CONVERT(nvarchar(50), {qident(type_col)}) = ?")
        params.append(type)

    if clean_keyword:
        where.append(f"{option_expression} LIKE ? ESCAPE N'~'")
        params.append(escape_like_prefix(clean_keyword) + "%")

    sql = (
        f"SELECT DISTINCT TOP ({limit}) "
        f"{option_expression} AS [OptionValue] "
        f"FROM {source_sql} "
        f"WHERE {' AND '.join(where)} "
        f"ORDER BY {option_expression};"
    )

    try:
        with get_connection() as conn:
            rows = conn.cursor().execute(sql, *params).fetchall()
    except pyodbc.Error as exc:
        raise HTTPException(
            status_code=500,
            detail=f"讀取搜尋選項失敗：{exc}",
        ) from exc

    values = [
        str(row[0]).strip()
        for row in rows
        if row[0] is not None and str(row[0]).strip()
    ]

    if field == "greige_type":
        options = [
            {
                "value": value,
                "label": f"{value}－{GREIGE_TYPE_LABELS.get(value, value)}",
            }
            for value in values
        ]
    elif field == "zero_project":
        zero_labels = {"Y": "Y－是", "N": "N－否"}
        options = [
            {
                "value": value,
                "label": zero_labels.get(value.upper(), value),
            }
            for value in values
        ]
    else:
        options = [{"value": value, "label": value} for value in values]

    return {
        "type": type,
        "field": field,
        "keyword": clean_keyword,
        "options": options,
    }


@app.get("/api/detail/options")
def detail_options(type: Literal["ALL", "1", "2", "3", "4", "8", "9", "D", "E", "F"] = Query(...)) -> dict[str, Any]:
    view_name = DETAIL_VIEW_MAP[type]
    source_sql, all_columns = detail_source_sql(view_name)
    mapping = resolve_detail_mapping(view_name, all_columns)
    option_limit = int(os.getenv("OPTION_LIMIT", "500"))

    result: dict[str, list[str]] = {}
    for key in (
        "item_no",
        "warehouse",
        "color_no",
        "zero_project",
        "salesperson",
        "business_department",
        "grade",
        "reason_code",
        "customer",
        "order_type",
        "greige_type",
    ):
        if key == "greige_type" and type != "2":
            result[key] = []
            continue
        column = mapping.get(key)
        if not column:
            result[key] = []
            continue

        where = [f"{qident(column)} IS NOT NULL"]
        params: list[Any] = []

        # 舊版 options API 亦套用固定排除規則。
        add_config_exclusion_filters(where, params, mapping)

        if type in SHARED_TYPE_CODES:
            type_col = require_field(mapping, "type_discriminator", "類型")
            where.append(f"CONVERT(nvarchar(50), {qident(type_col)}) = ?")
            params.append(type)

        legacy_expression = (
            f"LTRIM(RTRIM(CONVERT(nvarchar(4000), {qident(column)})))"
        )
        if key == "color_no":
            legacy_expression += " COLLATE Latin1_General_100_BIN2"

        sql = (
            f"SELECT DISTINCT TOP ({option_limit}) "
            f"{legacy_expression} AS [OptionValue] "
            f"FROM {source_sql} "
            f"WHERE {' AND '.join(where)} "
            f"ORDER BY {legacy_expression};"
        )
        with get_connection() as conn:
            rows = conn.cursor().execute(sql, *params).fetchall()
        result[key] = [
            str(row[0]).strip()
            for row in rows
            if row[0] is not None and str(row[0]).strip()
        ]

    return {
        "type": type,
        "view_name": view_name,
        "options": result,
        "mapping": mapping,
    }


@app.post("/api/detail/search")
def detail_search(req: DetailSearchRequest) -> dict[str, Any]:
    return fetch_detail_rows(req, export=False)


def workbook_response(
    title: str,
    columns: list[dict[str, Any]],
    rows: list[dict[str, Any]],
    filename: str,
) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    column_names = [str(item["name"]) for item in columns]
    column_labels = [
        str(item.get("label") or item["name"])
        for item in columns
    ]
    ws.append(column_labels)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(col) for col in column_names])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for index, (name, label) in enumerate(
        zip(column_names, column_labels),
        start=1,
    ):
        sample_lengths = [len(str(label))]
        for row in rows[:200]:
            value = row.get(name)
            if value is not None:
                sample_lengths.append(len(str(value)))
        ws.column_dimensions[get_column_letter(index)].width = min(
            max(sample_lengths) + 2,
            40,
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def fast_detail_xlsx_response(
    req: DetailSearchRequest,
) -> Response:
    """
    大量明細 Excel 快速匯出（Excel 相容版）。

    原則：
    - SQL 不先 COUNT_BIG。
    - SQL 分批 fetchmany，不 fetchall。
    - openpyxl write_only 降低記憶體使用量。
    - 未指定欄位排序時不做 ORDER BY。
    - 不在 write_only 完成後追加 AutoFilter / 欄寬等可能造成
      舊版 Excel 相容性問題的工作表結構。
    - 產出後先驗證 OOXML ZIP 結構，再以一般 Response 一次回傳，
      避免 XLSX 以 chunked StreamingResponse 傳輸時被環境干擾。
    """

    validate_paging(req)

    view_name = DETAIL_VIEW_MAP[req.inventory_type]
    source_sql, all_columns = detail_source_sql(view_name)
    metadata = visible_metadata(view_name)
    visible_columns = [item["name"] for item in metadata]
    mapping = resolve_detail_mapping(view_name, all_columns)
    where_sql, params = build_detail_filter_sql(req, mapping)

    if req.sort_field and req.sort_field not in visible_columns:
        raise HTTPException(
            status_code=400,
            detail="排序欄位不屬於目前顯示欄位。",
        )

    max_export_rows = int(os.getenv("MAX_EXPORT_ROWS", "100000"))
    select_list = build_detail_select_list(visible_columns, mapping)

    order_sql = ""
    if req.sort_field:
        order_sql = (
            f" ORDER BY {qident(req.sort_field)} {req.sort_direction}"
        )

    sql = (
        f"SELECT TOP ({max_export_rows}) {select_list} "
        f"FROM {source_sql} "
        f"WHERE {where_sql}"
        f"{order_sql};"
    )

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="庫存明細")

    # write_only 模式下，需要在第一列資料寫入前設定凍結窗格。
    ws.freeze_panes = "A2"

    column_labels = [
        str(item.get("label") or item["name"])
        for item in metadata
    ]

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_row = []
    for label in column_labels:
        cell = WriteOnlyCell(ws, value=label)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        header_row.append(cell)
    ws.append(header_row)

    type_column = mapping.get("type_discriminator")
    greige_type_column = mapping.get("greige_type")

    type_index = (
        visible_columns.index(type_column)
        if type_column in visible_columns
        else None
    )
    greige_type_index = (
        visible_columns.index(greige_type_column)
        if greige_type_column in visible_columns
        else None
    )

    fetch_size = max(
        500,
        int(os.getenv("EXPORT_FETCH_SIZE", "2000")),
    )

    exported_rows = 0

    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.arraysize = fetch_size
            cursor.execute(sql, *params)

            while True:
                batch = cursor.fetchmany(fetch_size)
                if not batch:
                    break

                for db_row in batch:
                    values = [
                        excel_cell_value(value)
                        for value in db_row
                    ]

                    if (
                        type_index is not None
                        and values[type_index] is not None
                    ):
                        raw_type = values[type_index]
                        type_code = str(raw_type).strip().upper()
                        values[type_index] = TYPE_LABELS.get(
                            type_code,
                            raw_type,
                        )

                    if (
                        greige_type_index is not None
                        and values[greige_type_index] is not None
                    ):
                        raw_greige = values[greige_type_index]
                        greige_code = str(raw_greige).strip()
                        values[greige_type_index] = (
                            GREIGE_TYPE_LABELS.get(
                                greige_code,
                                raw_greige,
                            )
                        )

                    ws.append(values)
                    exported_rows += 1

                    if exported_rows >= max_export_rows:
                        break

                if exported_rows >= max_export_rows:
                    break

    except pyodbc.Error as exc:
        raise HTTPException(
            status_code=500,
            detail=f"SQL Server 匯出失敗：{exc}",
        ) from exc

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()

    # XLSX 必須是合法 OOXML ZIP；送出前先做最基本完整性檢查。
    if not content.startswith(b"PK"):
        raise HTTPException(
            status_code=500,
            detail="Excel 產出失敗：檔案不是有效的 XLSX 格式。",
        )

    verify_stream = io.BytesIO(content)
    if not zipfile.is_zipfile(verify_stream):
        raise HTTPException(
            status_code=500,
            detail="Excel 產出失敗：XLSX ZIP 結構不完整。",
        )

    verify_stream.seek(0)
    try:
        with zipfile.ZipFile(verify_stream, "r") as zf:
            required_parts = {
                "[Content_Types].xml",
                "xl/workbook.xml",
                "xl/worksheets/sheet1.xml",
            }
            missing_parts = required_parts.difference(zf.namelist())
            if missing_parts:
                raise HTTPException(
                    status_code=500,
                    detail=(
                        "Excel 產出失敗：缺少必要 XLSX 結構："
                        + "、".join(sorted(missing_parts))
                    ),
                )
    except zipfile.BadZipFile as exc:
        raise HTTPException(
            status_code=500,
            detail="Excel 產出失敗：XLSX ZIP 已損毀。",
        ) from exc

    filename = (
        f"inventory_detail_type_{req.inventory_type}_"
        f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )

    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            ),
            "X-Exported-Rows": str(exported_rows),
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/detail/export/xlsx")
def detail_export_xlsx(
    req: DetailSearchRequest,
) -> Response:
    return fast_detail_xlsx_response(req)


@app.post("/api/detail/export/csv")
def detail_export_csv(req: DetailSearchRequest) -> StreamingResponse:
    result = fetch_detail_rows(req, export=True)
    columns = result["columns"]
    column_names = [str(item["name"]) for item in columns]
    column_labels = [
        str(item.get("label") or item["name"])
        for item in columns
    ]

    text = io.StringIO(newline="")
    writer = csv.writer(text)
    writer.writerow(column_labels)
    for row in result["rows"]:
        writer.writerow([row.get(name) for name in column_names])

    content = ("\ufeff" + text.getvalue()).encode("utf-8")
    filename = f"inventory_detail_type_{req.inventory_type}_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
