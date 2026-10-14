from __future__ import annotations

import io
import os
import re
from datetime import datetime
from urllib.parse import quote
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from db import get_connection

router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])

# -----------------------------------------------------------------------------
# CXCR002 dashboard source
# -----------------------------------------------------------------------------
_TABLE_NAME_RE = re.compile(r"^[A-Za-z0-9_]+(?:\.[A-Za-z0-9_]+){1,2}$")
CXCR002_TABLE = os.getenv("CXCR002_TABLE", "MGT_MRT.dbo.CXCR002_Fact").strip()
CXCR002_OPERATING_CENTER = os.getenv(
    "CXCR002_OPERATING_CENTER", "SNX_202012"
).strip()
# CXCR002 / cxcr002.4gl current production context uses cma07 = 3.
# Set this env var to blank if a site intentionally wants all categories.
CXCR002_INVENTORY_CATEGORY_CODE = os.getenv(
    "CXCR002_INVENTORY_CATEGORY_CODE", "3"
).strip()

if not _TABLE_NAME_RE.fullmatch(CXCR002_TABLE):
    raise RuntimeError(
        "CXCR002_TABLE must be a 2- or 3-part SQL Server identifier, "
        "for example MGT_MRT.dbo.CXCR002_Fact."
    )


def _qtable(name: str) -> str:
    return ".".join(f"[{part}]" for part in name.split("."))


FACT = _qtable(CXCR002_TABLE)

TYPE_LABELS = {"1": "布", "2": "胚布", "3": "紗"}
TOP_TYPE_LABELS = {"1": "布", "2": "胚布", "3": "紗"}
# 「布」同時涵蓋 ItemTypeCode 1 / D；頁面顯示仍統一為「布」。
TOP_TYPE_CODES = {"1": ["1", "D"], "2": ["2"], "3": ["3"]}
FABRIC_LABELS = {"1": "針織", "2": "平織"}
ELASTIC_LABELS = {"1": "彈性", "2": "非彈性"}

# CXCR002 current age-band boundaries.  These match the 2-year/3-year report
# semantics and avoid using CurrentStockQuantity, which is a current/live value.
AGE_BUCKETS = [
    (1, "0-30天", 0, 30),
    (2, "31-90天", 31, 90),
    (3, "91-180天", 91, 180),
    (4, "181-365天", 181, 365),
    (5, "1年-1.5年", 366, 548),
    (6, "1.5年-2年", 549, 730),
    (7, "2年-3年", 731, 1095),
    (8, "3年-5年", 1096, 1825),
    (9, "5年以上", 1826, None),
]


class DashboardRequest(BaseModel):
    # Names retained for front-end compatibility. Values are CXCR002 ReportPeriod
    # (YYYYMM), not InventoryInquiry snapshot dates.
    snapshot_date: str | None = None
    base_snapshot_date: str | None = None
    compare_snapshot_date: str | None = None
    inventory_types: list[str] = Field(default_factory=list)
    fabric_types: list[str] = Field(default_factory=list)
    elastic_flags: list[str] = Field(default_factory=list)
    project_flags: list[str] = Field(default_factory=list)
    # None means「全部」; numeric values keep the existing Top-N behavior.
    top_n: int | None = 20
    # Page 4 Top-N dedicated filters.
    # top_type: 1=布, 2=胚布, 3=紗
    # top_age_to=None means no upper bound (NULL / 以上).
    top_type: str = "2"
    top_age_from: float | None = 0
    top_age_to: float | None = 2
    # Optional fine-grained buckets when From=0 and To=1.
    # Supports multi-select: all / 0_30 / 31_90 / 91_180 / 181_365
    # String is still accepted for backward compatibility with older front-end builds.
    top_age_subrange: list[str] | str | None = Field(default_factory=lambda: ["all"])
    # Backward compatibility with the immediately previous Page 4 patch.
    top_category: str | None = None
    top_age_range: str | None = None


def _as_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _pct(numerator: float, denominator: float) -> float:
    return 0.0 if denominator == 0 else numerator * 100.0 / denominator


def _provision_rate(provision: float, amount: float) -> float:
    # Procurement meeting display treats provision as a deduction, hence negative.
    return 0.0 if amount == 0 else -provision * 100.0 / amount


def _query(sql: str, params: list[Any] | tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, *params)
            if cursor.description is None:
                return []
            columns = [item[0] for item in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]
    except Exception as exc:  # pyodbc is deliberately kept inside db.py
        raise HTTPException(status_code=500, detail=f"CXCR002 查詢失敗：{exc}") from exc


def _column(name: str, label: str | None = None) -> dict[str, str]:
    return {"name": name, "label": label or name}


def _validate_period(value: str | int | None, label: str) -> int:
    try:
        period = int(str(value).strip())
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{label} 必須是 YYYYMM。")
    month = period % 100
    if period < 190001 or period > 299912 or not 1 <= month <= 12:
        raise HTTPException(status_code=400, detail=f"{label} 格式錯誤：{period}。")
    return period


def _periods() -> list[int]:
    where, params = _fixed_where()
    rows = _query(
        f"""
        SELECT TOP (36) CONVERT(int, ReportPeriod) AS ReportPeriod
        FROM {FACT}
        WHERE {' AND '.join(where)}
          AND TRY_CONVERT(int, ReportPeriod) BETWEEN 190001 AND 299912
          AND TRY_CONVERT(int, ReportPeriod) % 100 BETWEEN 1 AND 12
        GROUP BY CONVERT(int, ReportPeriod)
        ORDER BY CONVERT(int, ReportPeriod) DESC;
        """,
        params,
    )
    return [int(row["ReportPeriod"]) for row in rows]


def _latest_period() -> int:
    periods = _periods()
    if not periods:
        raise HTTPException(
            status_code=500,
            detail=(
                f"{CXCR002_TABLE} 找不到 {CXCR002_OPERATING_CENTER} "
                "可使用的 ReportPeriod。"
            ),
        )
    return periods[0]


def _fixed_where() -> tuple[list[str], list[Any]]:
    where = ["OperatingCenterID = ?"]
    params: list[Any] = [CXCR002_OPERATING_CENTER]
    if CXCR002_INVENTORY_CATEGORY_CODE:
        where.append("CONVERT(nvarchar(20), InventoryCategoryCode) = ?")
        params.append(CXCR002_INVENTORY_CATEGORY_CODE)
    return where, params


def _clean_values(values: list[str], allowed: set[str] | None = None) -> list[str]:
    clean = []
    for value in values or []:
        text = str(value).strip()
        if not text:
            continue
        if allowed is not None and text not in allowed:
            continue
        if text not in clean:
            clean.append(text)
    return clean


def _dashboard_where(
    req: DashboardRequest,
    *,
    periods: list[int] | None = None,
    force_types: list[str] | None = None,
    force_fabric: str | None = None,
    three_years: bool = False,
    apply_structure_filters: bool = True,
) -> tuple[list[str], list[Any]]:
    where, params = _fixed_where()

    if periods:
        placeholders = ",".join("?" for _ in periods)
        where.append(f"CONVERT(int, ReportPeriod) IN ({placeholders})")
        params.extend(periods)

    if force_types is None:
        selected_types = _clean_values(req.inventory_types, {"2", "3"})
        selected_types = selected_types or ["2", "3"]
    else:
        # Forced report types may include finished fabric (1) for Page 4 Top-N.
        # The regular structure pages still expose only yarn/greige via their UI.
        selected_types = _clean_values(force_types, {"1", "2", "3", "D"})

    if selected_types:
        placeholders = ",".join("?" for _ in selected_types)
        where.append(f"CONVERT(nvarchar(20), ItemTypeCode) IN ({placeholders})")
        params.extend(selected_types)

    if force_fabric:
        where.append("CONVERT(nvarchar(20), FabricTypeCode) = ?")
        params.append(force_fabric)
    elif apply_structure_filters:
        fabrics = _clean_values(req.fabric_types, {"1", "2"})
        if fabrics:
            placeholders = ",".join("?" for _ in fabrics)
            where.append(f"CONVERT(nvarchar(20), FabricTypeCode) IN ({placeholders})")
            params.extend(fabrics)

    elastics = _clean_values(req.elastic_flags, {"1", "2"})
    if elastics:
        placeholders = ",".join("?" for _ in elastics)
        where.append(f"CONVERT(nvarchar(20), ElasticFlag) IN ({placeholders})")
        params.extend(elastics)

    projects = _clean_values([str(v).upper() for v in req.project_flags], {"Y", "N"})
    if projects:
        placeholders = ",".join("?" for _ in projects)
        where.append(
            "CASE WHEN UPPER(LTRIM(RTRIM(ISNULL(CONVERT(nvarchar(10), "
            f"Project302Flag), N'')))) = N'Y' THEN N'Y' ELSE N'N' END IN ({placeholders})"
        )
        params.extend(projects)

    if three_years:
        where.append("AgeDays BETWEEN 0 AND 1095")
    else:
        where.append("AgeDays >= 0")

    return where, params


def _age_case_sql() -> tuple[str, str]:
    parts_name: list[str] = []
    parts_sort: list[str] = []
    for sort, label, low, high in AGE_BUCKETS:
        if high is None:
            cond = f"AgeDays >= {low}"
        else:
            cond = f"AgeDays BETWEEN {low} AND {high}"
        parts_name.append(f"WHEN {cond} THEN N'{label}'")
        parts_sort.append(f"WHEN {cond} THEN {sort}")
    return (
        "CASE " + " ".join(parts_name) + " ELSE N'其他' END",
        "CASE " + " ".join(parts_sort) + " ELSE 99 END",
    )


@router.get("/options")
def dashboard_options() -> dict[str, Any]:
    periods = _periods()
    latest = periods[0] if periods else None
    previous = periods[1] if len(periods) > 1 else None

    return {
        "analysis_source": CXCR002_TABLE,
        "data_scope": "CXCR002 月結庫存 Fact",
        "latest_report_period": latest,
        "previous_report_period": previous,
        # Keep existing response names so current JS can be upgraded without
        # breaking old callers. These are YYYYMM ReportPeriod values.
        "snapshot_dates": [str(value) for value in periods],
        "snapshot_enabled": len(periods) >= 2,
        "latest_snapshot_date": str(latest) if latest else None,
        "previous_snapshot_date": str(previous) if previous else None,
        "mapping": {
            "report_period": "ReportPeriod",
            "inventory_type": "ItemTypeCode / TypeClassificationName",
            "fabric_type": "FabricTypeCode",
            "elastic_flag": "ElasticFlag",
            "project_302": "Project302Flag",
            "quantity": "InventoryQuantity",
            "inventory_amount": "InventoryCostAmountNTD",
            "book_value": "BookValueAmountNTD",
            "provision": "ProvisionAmountNTD",
        },
        "provision_mode": "CXCR002 ProvisionAmountNTD / BookValueAmountNTD",
        "options": {
            "inventory_type": [
                {"value": "3", "label": "紗"},
                {"value": "2", "label": "胚布"},
            ],
            "fabric_type": [
                {"value": "1", "label": "針織"},
                {"value": "2", "label": "平織"},
            ],
            "elastic_flag": [
                {"value": "1", "label": "彈性"},
                {"value": "2", "label": "非彈性"},
            ],
            "project_302": [
                {"value": "Y", "label": "Y－是"},
                {"value": "N", "label": "N－否"},
            ],
        },
        "available_filters": {
            "inventory_type": True,
            "fabric_type": True,
            "elastic_flag": True,
            "project_302": True,
        },
    }


@router.post("/aging")
def aging(req: DashboardRequest) -> dict[str, Any]:
    period = _latest_period()
    where, params = _dashboard_where(req, periods=[period])

    bucket_selects: list[str] = []
    for _, label, low, high in AGE_BUCKETS:
        if high is None:
            cond = f"AgeDays >= {low}"
        else:
            cond = f"AgeDays BETWEEN {low} AND {high}"
        bucket_selects.append(
            f"SUM(CASE WHEN {cond} THEN ISNULL(InventoryCostAmountNTD,0) ELSE 0 END) AS [{label}]"
        )

    rows = _query(
        f"""
        SELECT
            CASE WHEN UPPER(LTRIM(RTRIM(ISNULL(CONVERT(nvarchar(10), Project302Flag),N''))))=N'Y'
                 THEN N'Y' ELSE N'N' END AS Project302,
            CONVERT(nvarchar(20), ItemTypeCode) AS ItemTypeCode,
            CONVERT(nvarchar(20), FabricTypeCode) AS FabricTypeCode,
            CONVERT(nvarchar(20), ElasticFlag) AS ElasticFlag,
            {', '.join(bucket_selects)}
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY
            CASE WHEN UPPER(LTRIM(RTRIM(ISNULL(CONVERT(nvarchar(10), Project302Flag),N''))))=N'Y'
                 THEN N'Y' ELSE N'N' END,
            CONVERT(nvarchar(20), ItemTypeCode),
            CONVERT(nvarchar(20), FabricTypeCode),
            CONVERT(nvarchar(20), ElasticFlag)
        ORDER BY Project302, ItemTypeCode, FabricTypeCode, ElasticFlag;
        """,
        params,
    )

    bucket_totals = {label: 0.0 for _, label, _, _ in AGE_BUCKETS}
    for source in rows:
        for _, label, _, _ in AGE_BUCKETS:
            bucket_totals[label] += _as_number(source.get(label))
    grand_total = sum(bucket_totals.values())

    output: list[dict[str, Any]] = []
    for source in rows:
        type_code = str(source.get("ItemTypeCode") or "").strip()
        row: dict[str, Any] = {
            "302專案否": source.get("Project302") or "N",
            "類型分類": TYPE_LABELS.get(type_code, type_code),
            "布種別": (
                FABRIC_LABELS.get(str(source.get("FabricTypeCode") or "").strip(), "--")
                if type_code == "2"
                else "--"
            ),
            "彈性否": (
                ELASTIC_LABELS.get(str(source.get("ElasticFlag") or "").strip(), "--")
                if type_code == "2"
                else "--"
            ),
        }
        row_total = 0.0
        for _, label, _, _ in AGE_BUCKETS:
            amount = _as_number(source.get(label))
            row[f"{label}存貨金額"] = amount
            row[f"{label}百分比"] = _pct(amount, bucket_totals[label])
            row_total += amount
        row["存貨金額"] = row_total
        row["總百分比"] = _pct(row_total, grand_total)
        output.append(row)

    total_row: dict[str, Any] = {
        "302專案否": "",
        "類型分類": "總計",
        "布種別": "",
        "彈性否": "",
    }
    for _, label, _, _ in AGE_BUCKETS:
        total_row[f"{label}存貨金額"] = bucket_totals[label]
        total_row[f"{label}百分比"] = 100.0 if bucket_totals[label] else 0.0
    total_row["存貨金額"] = grand_total
    total_row["總百分比"] = 100.0 if grand_total else 0.0
    output.append(total_row)

    columns = [
        _column("302專案否"),
        _column("類型分類"),
        _column("布種別"),
        _column("彈性否"),
    ]
    for _, label, _, _ in AGE_BUCKETS:
        columns.extend([
            _column(f"{label}存貨金額"),
            _column(f"{label}百分比"),
        ])
    columns.extend([_column("存貨金額"), _column("總百分比")])

    return {
        "report_period": period,
        "columns": columns,
        "rows": output,
    }


@router.post("/category")
def category(req: DashboardRequest) -> dict[str, Any]:
    period = _latest_period()
    where, params = _dashboard_where(req, periods=[period])
    age_group = """
        CASE
            WHEN AgeDays BETWEEN 0 AND 365 THEN N'1年以內'
            WHEN AgeDays BETWEEN 366 AND 730 THEN N'1-2年'
            WHEN AgeDays BETWEEN 731 AND 1095 THEN N'2-3年'
            ELSE N'3年以上'
        END
    """
    age_sort = """
        CASE
            WHEN AgeDays BETWEEN 0 AND 365 THEN 1
            WHEN AgeDays BETWEEN 366 AND 730 THEN 2
            WHEN AgeDays BETWEEN 731 AND 1095 THEN 3
            ELSE 4
        END
    """
    source = _query(
        f"""
        SELECT
            {age_group} AS AgeGroup,
            {age_sort} AS AgeSort,
            CONVERT(nvarchar(20), ItemTypeCode) AS ItemTypeCode,
            SUM(ISNULL(InventoryCostAmountNTD,0)) AS Amount
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY {age_group}, {age_sort}, CONVERT(nvarchar(20), ItemTypeCode)
        ORDER BY AgeSort, ItemTypeCode;
        """,
        params,
    )

    values: dict[tuple[str, str], float] = defaultdict(float)
    type_totals = {"2": 0.0, "3": 0.0}
    for row in source:
        age = str(row.get("AgeGroup") or "")
        code = str(row.get("ItemTypeCode") or "")
        amount = _as_number(row.get("Amount"))
        values[(age, code)] += amount
        if code in type_totals:
            type_totals[code] += amount
    combined_total = type_totals["2"] + type_totals["3"]

    output: list[dict[str, Any]] = []
    for age in ["1年以內", "1-2年", "2-3年", "3年以上"]:
        yarn = values[(age, "3")]
        greige = values[(age, "2")]
        combined = yarn + greige
        output.append({
            "庫齡區間": age,
            "紗存貨金額": yarn,
            "紗占比百分比": _pct(yarn, type_totals["3"]),
            "胚布存貨金額": greige,
            "胚布占比百分比": _pct(greige, type_totals["2"]),
            "合計存貨金額": combined,
            "合計占比百分比": _pct(combined, combined_total),
        })

    output.append({
        "庫齡區間": "總計",
        "紗存貨金額": type_totals["3"],
        "紗占比百分比": 100.0 if type_totals["3"] else 0.0,
        "胚布存貨金額": type_totals["2"],
        "胚布占比百分比": 100.0 if type_totals["2"] else 0.0,
        "合計存貨金額": combined_total,
        "合計占比百分比": 100.0 if combined_total else 0.0,
    })

    columns = [
        _column("庫齡區間"),
        _column("紗存貨金額"),
        _column("紗占比百分比"),
        _column("胚布存貨金額"),
        _column("胚布占比百分比"),
        _column("合計存貨金額"),
        _column("合計占比百分比"),
    ]
    return {"report_period": period, "columns": columns, "rows": output}


def _comparison_source(req: DashboardRequest, a: int, b: int) -> list[dict[str, Any]]:
    where, params = _dashboard_where(req, periods=[a, b])
    age_name, age_sort = _age_case_sql()
    return _query(
        f"""
        SELECT
            CONVERT(int, ReportPeriod) AS ReportPeriod,
            CONVERT(nvarchar(20), ItemTypeCode) AS ItemTypeCode,
            CONVERT(nvarchar(20), FabricTypeCode) AS FabricTypeCode,
            {age_sort} AS AgeSort,
            {age_name} AS AgeBucket,
            SUM(ISNULL(InventoryQuantity,0)) AS Quantity,
            SUM(ISNULL(InventoryCostAmountNTD,0)) AS Amount,
            SUM(ISNULL(BookValueAmountNTD,0)) AS BookValue,
            SUM(ISNULL(ProvisionAmountNTD,0)) AS Provision
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY
            CONVERT(int, ReportPeriod),
            CONVERT(nvarchar(20), ItemTypeCode),
            CONVERT(nvarchar(20), FabricTypeCode),
            {age_sort}, {age_name}
        ORDER BY AgeSort, ItemTypeCode, FabricTypeCode, ReportPeriod;
        """,
        params,
    )


def _period_total(source: list[dict[str, Any]], period: int, type_code: str) -> float:
    return sum(
        _as_number(row.get("Amount"))
        for row in source
        if int(row.get("ReportPeriod") or 0) == period
        and str(row.get("ItemTypeCode") or "") == type_code
    )


@router.post("/period-comparison")
def period_comparison(req: DashboardRequest) -> dict[str, Any]:
    periods = _periods()
    if len(periods) < 2:
        return {
            "available": False,
            "message": "CXCR002_Fact 目前不足兩個 ReportPeriod，無法做兩期比較。",
        }

    a = _validate_period(req.base_snapshot_date or periods[1], "基準月份")
    b = _validate_period(req.compare_snapshot_date or periods[0], "比較月份")
    if a == b:
        raise HTTPException(status_code=400, detail="基準月份與比較月份不可相同。")
    if a not in periods or b not in periods:
        raise HTTPException(
            status_code=400,
            detail="指定月份不在 CXCR002_Fact 可用 ReportPeriod 清單中。",
        )

    source = _comparison_source(req, a, b)
    labels = [label for _, label, _, _ in AGE_BUCKETS]

    # ---- Yarn ---------------------------------------------------------------
    yarn_total_a = _period_total(source, a, "3")
    yarn_total_b = _period_total(source, b, "3")
    yarn_rows: list[dict[str, Any]] = []

    for label in labels:
        data: dict[int, dict[str, float]] = {}
        for period in (a, b):
            matches = [
                row for row in source
                if int(row.get("ReportPeriod") or 0) == period
                and str(row.get("ItemTypeCode") or "") == "3"
                and str(row.get("AgeBucket") or "") == label
            ]
            data[period] = {
                "qty": sum(_as_number(r.get("Quantity")) for r in matches),
                "amount": sum(_as_number(r.get("Amount")) for r in matches),
                "book": sum(_as_number(r.get("BookValue")) for r in matches),
                "provision": sum(_as_number(r.get("Provision")) for r in matches),
            }
        av, bv = data[a], data[b]
        yarn_rows.append({
            "庫齡區間": label,
            f"{a}庫存數量": av["qty"],
            f"{a}存貨金額": av["amount"],
            f"{a}金額百分比": _pct(av["amount"], yarn_total_a),
            f"{a}帳面價值": av["book"],
            f"{a}提列率": _provision_rate(av["provision"], av["amount"]),
            f"{b}庫存數量": bv["qty"],
            f"{b}存貨金額": bv["amount"],
            f"{b}金額百分比": _pct(bv["amount"], yarn_total_b),
            f"{b}帳面價值": bv["book"],
            f"{b}提列率": _provision_rate(bv["provision"], bv["amount"]),
            "差異庫存數量": av["qty"] - bv["qty"],
            "差異存貨金額": av["amount"] - bv["amount"],
        })

    def _sum_field(rows: list[dict[str, Any]], key: str) -> float:
        return sum(_as_number(row.get(key)) for row in rows)

    yarn_rows.append({
        "庫齡區間": "總計",
        f"{a}庫存數量": _sum_field(yarn_rows, f"{a}庫存數量"),
        f"{a}存貨金額": yarn_total_a,
        f"{a}金額百分比": 100.0 if yarn_total_a else 0.0,
        f"{a}帳面價值": _sum_field(yarn_rows, f"{a}帳面價值"),
        f"{a}提列率": _provision_rate(
            sum(
                _as_number(r.get("Provision"))
                for r in source
                if int(r.get("ReportPeriod") or 0) == a
                and str(r.get("ItemTypeCode") or "") == "3"
            ),
            yarn_total_a,
        ),
        f"{b}庫存數量": _sum_field(yarn_rows, f"{b}庫存數量"),
        f"{b}存貨金額": yarn_total_b,
        f"{b}金額百分比": 100.0 if yarn_total_b else 0.0,
        f"{b}帳面價值": _sum_field(yarn_rows, f"{b}帳面價值"),
        f"{b}提列率": _provision_rate(
            sum(
                _as_number(r.get("Provision"))
                for r in source
                if int(r.get("ReportPeriod") or 0) == b
                and str(r.get("ItemTypeCode") or "") == "3"
            ),
            yarn_total_b,
        ),
        "差異庫存數量": _sum_field(yarn_rows, f"{a}庫存數量") - _sum_field(yarn_rows, f"{b}庫存數量"),
        "差異存貨金額": yarn_total_a - yarn_total_b,
    })

    yarn_columns = [_column("庫齡區間")]
    for period in (a, b):
        yarn_columns.extend([
            _column(f"{period}庫存數量"),
            _column(f"{period}存貨金額"),
            _column(f"{period}金額百分比"),
            _column(f"{period}帳面價值"),
            _column(f"{period}提列率"),
        ])
    yarn_columns.extend([_column("差異庫存數量"), _column("差異存貨金額")])

    # ---- Greige -------------------------------------------------------------
    greige_total_a = _period_total(source, a, "2")
    greige_total_b = _period_total(source, b, "2")
    greige_rows: list[dict[str, Any]] = []
    for label in labels:
        data: dict[int, dict[str, float]] = {}
        for period in (a, b):
            matches = [
                row for row in source
                if int(row.get("ReportPeriod") or 0) == period
                and str(row.get("ItemTypeCode") or "") == "2"
                and str(row.get("AgeBucket") or "") == label
            ]
            data[period] = {
                "flat_qty": sum(
                    _as_number(r.get("Quantity"))
                    for r in matches if str(r.get("FabricTypeCode") or "") == "2"
                ),
                "knit_qty": sum(
                    _as_number(r.get("Quantity"))
                    for r in matches if str(r.get("FabricTypeCode") or "") == "1"
                ),
                "amount": sum(_as_number(r.get("Amount")) for r in matches),
                "book": sum(_as_number(r.get("BookValue")) for r in matches),
                "provision": sum(_as_number(r.get("Provision")) for r in matches),
            }
        av, bv = data[a], data[b]
        greige_rows.append({
            "庫齡區間": label,
            f"{a}平織數量": av["flat_qty"],
            f"{a}針織數量": av["knit_qty"],
            f"{a}存貨金額": av["amount"],
            f"{a}金額百分比": _pct(av["amount"], greige_total_a),
            f"{a}帳面價值": av["book"],
            f"{a}提列率": _provision_rate(av["provision"], av["amount"]),
            f"{b}平織數量": bv["flat_qty"],
            f"{b}針織數量": bv["knit_qty"],
            f"{b}存貨金額": bv["amount"],
            f"{b}金額百分比": _pct(bv["amount"], greige_total_b),
            f"{b}帳面價值": bv["book"],
            f"{b}提列率": _provision_rate(bv["provision"], bv["amount"]),
            "差異平織數量": av["flat_qty"] - bv["flat_qty"],
            "差異針織數量": av["knit_qty"] - bv["knit_qty"],
            "差異存貨金額": av["amount"] - bv["amount"],
        })

    greige_rows.append({
        "庫齡區間": "總計",
        f"{a}平織數量": _sum_field(greige_rows, f"{a}平織數量"),
        f"{a}針織數量": _sum_field(greige_rows, f"{a}針織數量"),
        f"{a}存貨金額": greige_total_a,
        f"{a}金額百分比": 100.0 if greige_total_a else 0.0,
        f"{a}帳面價值": _sum_field(greige_rows, f"{a}帳面價值"),
        f"{a}提列率": _provision_rate(
            sum(
                _as_number(r.get("Provision"))
                for r in source
                if int(r.get("ReportPeriod") or 0) == a
                and str(r.get("ItemTypeCode") or "") == "2"
            ),
            greige_total_a,
        ),
        f"{b}平織數量": _sum_field(greige_rows, f"{b}平織數量"),
        f"{b}針織數量": _sum_field(greige_rows, f"{b}針織數量"),
        f"{b}存貨金額": greige_total_b,
        f"{b}金額百分比": 100.0 if greige_total_b else 0.0,
        f"{b}帳面價值": _sum_field(greige_rows, f"{b}帳面價值"),
        f"{b}提列率": _provision_rate(
            sum(
                _as_number(r.get("Provision"))
                for r in source
                if int(r.get("ReportPeriod") or 0) == b
                and str(r.get("ItemTypeCode") or "") == "2"
            ),
            greige_total_b,
        ),
        "差異平織數量": _sum_field(greige_rows, f"{a}平織數量") - _sum_field(greige_rows, f"{b}平織數量"),
        "差異針織數量": _sum_field(greige_rows, f"{a}針織數量") - _sum_field(greige_rows, f"{b}針織數量"),
        "差異存貨金額": greige_total_a - greige_total_b,
    })

    greige_columns = [_column("庫齡區間")]
    for period in (a, b):
        greige_columns.extend([
            _column(f"{period}平織數量"),
            _column(f"{period}針織數量"),
            _column(f"{period}存貨金額"),
            _column(f"{period}金額百分比"),
            _column(f"{period}帳面價值"),
            _column(f"{period}提列率"),
        ])
    greige_columns.extend([
        _column("差異平織數量"),
        _column("差異針織數量"),
        _column("差異存貨金額"),
    ])

    return {
        "available": True,
        "base_snapshot_date": str(a),
        "compare_snapshot_date": str(b),
        "base_report_period": a,
        "compare_report_period": b,
        "difference_definition": "基準月份－比較月份",
        "yarn": {"columns": yarn_columns, "rows": yarn_rows},
        # Keep key 'knit' for compatibility with the current JS; this payload is
        # actually the whole greige table (flat + knit quantity split).
        "knit": {"columns": greige_columns, "rows": greige_rows},
        "greige": {"columns": greige_columns, "rows": greige_rows},
    }


def _normalize_top_type(req: DashboardRequest) -> str:
    # Prefer the new field; retain previous top_category as fallback.
    code = str(req.top_type or req.top_category or "2").strip().upper()
    return code if code in TOP_TYPE_LABELS else "2"


def _year_boundary_days(value: float | int | Decimal) -> int:
    """Convert report-year boundary to the same day convention used by CXCR002.

    Examples: 1 -> 365, 1.5 -> 548, 2 -> 730, 3 -> 1095.
    """
    years = Decimal(str(value))
    return int((years * Decimal("365")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _top_age_filter(
    age_from: float | None,
    age_to: float | None,
    age_subrange: list[str] | str | None = None,
) -> tuple[str, list[Any], float, float | None, str]:
    try:
        start_year = float(0 if age_from is None else age_from)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="庫齡 From 必須是非負數。")

    if start_year < 0:
        raise HTTPException(status_code=400, detail="庫齡 From 不可小於 0。")

    end_year: float | None
    if age_to is None:
        end_year = None
    else:
        try:
            end_year = float(age_to)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="庫齡 To 必須是非負數或留空。")
        if end_year < 0:
            raise HTTPException(status_code=400, detail="庫齡 To 不可小於 0。")
        if end_year <= start_year:
            raise HTTPException(status_code=400, detail="庫齡 To 必須大於 From；若要查 From 以上，請將 To 留空。")

    # Special drill-down for 0 -> 1 year. Multiple first-year buckets can be
    # combined freely. SQL uses OR so non-contiguous selections are preserved.
    allowed_subranges = {
        "0_30": (0, 30, "0～30天"),
        "31_90": (31, 90, "31～90天"),
        "91_180": (91, 180, "91～180天"),
        "181_365": (181, 365, "181～365天"),
    }

    if isinstance(age_subrange, (list, tuple, set)):
        selected_subranges = [str(value or "").strip().lower() for value in age_subrange]
    elif age_subrange is None:
        selected_subranges = ["all"]
    else:
        selected_subranges = [str(age_subrange).strip().lower()]

    selected_subranges = [value for value in selected_subranges if value]
    if not selected_subranges:
        selected_subranges = ["all"]

    if start_year == 0 and end_year == 1:
        if "all" in selected_subranges:
            return "AgeDays BETWEEN ? AND ?", [0, 365], start_year, end_year, "0～1年（全部）"

        invalid = [value for value in selected_subranges if value not in allowed_subranges]
        if invalid:
            raise HTTPException(status_code=400, detail="一年內區間值不正確。")

        # Preserve the business bucket order regardless of click order.
        bucket_order = ["0_30", "31_90", "91_180", "181_365"]
        chosen = [value for value in bucket_order if value in set(selected_subranges)]
        if not chosen:
            return "AgeDays BETWEEN ? AND ?", [0, 365], start_year, end_year, "0～1年（全部）"

        clauses: list[str] = []
        params: list[Any] = []
        labels: list[str] = []
        for key in chosen:
            low_day, high_day, label = allowed_subranges[key]
            clauses.append("AgeDays BETWEEN ? AND ?")
            params.extend([low_day, high_day])
            labels.append(label)

        # If all four buckets are selected, simplify to the full first year.
        if len(chosen) == 4:
            return "AgeDays BETWEEN ? AND ?", [0, 365], start_year, end_year, "0～1年（全部）"

        return "(" + " OR ".join(clauses) + ")", params, start_year, end_year, "、".join(labels)

    # From=0 means include AgeDays=0. Any positive From is an exclusive year
    # boundary: From=2 starts at day 731, matching the established 2~3 year band.
    low_day = 0 if start_year == 0 else _year_boundary_days(start_year) + 1

    if end_year is None:
        age_sql = "AgeDays >= ?"
        params = [low_day]
        label = f"{start_year:g}年以上"
    else:
        high_day = _year_boundary_days(end_year)
        age_sql = "AgeDays BETWEEN ? AND ?"
        params = [low_day, high_day]
        if start_year == 0:
            label = f"0～{end_year:g}年"
        else:
            label = f"{start_year:g}～{end_year:g}年"

    return age_sql, params, start_year, end_year, label

def _format_age_span(min_days: Any, max_days: Any) -> str:
    try:
        low = int(min_days) if min_days is not None else None
        high = int(max_days) if max_days is not None else None
    except (TypeError, ValueError):
        return ""
    if low is None and high is None:
        return ""
    if low == high:
        return f"{low}天"
    if low is None:
        return f"～{high}天"
    if high is None:
        return f"{low}天以上"
    return f"{low}～{high}天"


def _top_category_source(
    req: DashboardRequest,
    period: int,
) -> tuple[list[dict[str, Any]], str, str, float, float | None, str]:
    type_code = _normalize_top_type(req)
    type_label = TOP_TYPE_LABELS[type_code]
    force_type_codes = TOP_TYPE_CODES[type_code]
    age_sql, age_params, age_from, age_to, age_label = _top_age_filter(
        req.top_age_from, req.top_age_to, req.top_age_subrange
    )

    where, params = _dashboard_where(
        req,
        periods=[period],
        force_types=force_type_codes,
        three_years=False,
        apply_structure_filters=False,
    )
    where.append(age_sql)
    params.extend(age_params)

    rows = _query(
        f"""
        SELECT
            CONVERT(nvarchar(200), ItemNo) AS ItemNo,
            CONVERT(nvarchar(40), InventoryUnitCode) AS InventoryUnitCode,
            MIN(AgeDays) AS MinAgeDays,
            MAX(AgeDays) AS MaxAgeDays,
            MIN(NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(20), FabricTypeCode))),N'')) AS MinFabricTypeCode,
            MAX(NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(20), FabricTypeCode))),N'')) AS MaxFabricTypeCode,
            SUM(ISNULL(InventoryQuantity,0)) AS Quantity,
            SUM(ISNULL(InventoryCostAmountNTD,0)) AS Amount,
            SUM(ISNULL(BookValueAmountNTD,0)) AS BookValue,
            SUM(ISNULL(ProvisionAmountNTD,0)) AS Provision
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY CONVERT(nvarchar(200), ItemNo), CONVERT(nvarchar(40), InventoryUnitCode)
        HAVING SUM(ISNULL(InventoryQuantity,0)) <> 0
            OR SUM(ISNULL(InventoryCostAmountNTD,0)) <> 0
        ORDER BY SUM(ISNULL(InventoryQuantity,0)) DESC, ItemNo;
        """,
        params,
    )
    # 布種別僅作為顯示欄位，不改變 Top-N 的料號排名粒度。
    # 正常情況同一料號會落在單一 FabricTypeCode；若來源真的混用，明確標示為「針織/平織」。
    for row in rows:
        min_fabric = str(row.get("MinFabricTypeCode") or "").strip()
        max_fabric = str(row.get("MaxFabricTypeCode") or "").strip()
        if min_fabric and min_fabric == max_fabric:
            row["FabricLabel"] = FABRIC_LABELS.get(min_fabric, min_fabric)
        elif {min_fabric, max_fabric} == {"1", "2"}:
            row["FabricLabel"] = "針織/平織"
        else:
            row["FabricLabel"] = FABRIC_LABELS.get(min_fabric, min_fabric) or FABRIC_LABELS.get(max_fabric, max_fabric)

    return rows, type_code, type_label, age_from, age_to, age_label


def _top_source(req: DashboardRequest, period: int, fabric_code: str) -> list[dict[str, Any]]:
    """Page 5 ranking source: greige only, but use Page 4's age From/To filter."""
    where, params = _dashboard_where(
        req,
        periods=[period],
        force_types=["2"],
        force_fabric=fabric_code,
        three_years=False,
        apply_structure_filters=False,
    )
    age_sql, age_params, _, _, _ = _top_age_filter(req.top_age_from, req.top_age_to, req.top_age_subrange)
    where.append(age_sql)
    params.extend(age_params)
    return _query(
        f"""
        SELECT
            CONVERT(nvarchar(200), ItemNo) AS ItemNo,
            CONVERT(nvarchar(40), InventoryUnitCode) AS InventoryUnitCode,
            MIN(AgeDays) AS MinAgeDays,
            MAX(AgeDays) AS MaxAgeDays,
            SUM(ISNULL(InventoryQuantity,0)) AS Quantity,
            SUM(ISNULL(InventoryCostAmountNTD,0)) AS Amount,
            SUM(ISNULL(BookValueAmountNTD,0)) AS BookValue,
            SUM(ISNULL(ProvisionAmountNTD,0)) AS Provision
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY CONVERT(nvarchar(200), ItemNo), CONVERT(nvarchar(40), InventoryUnitCode)
        HAVING SUM(ISNULL(InventoryQuantity,0)) <> 0
            OR SUM(ISNULL(InventoryCostAmountNTD,0)) <> 0
        ORDER BY SUM(ISNULL(InventoryQuantity,0)) DESC, ItemNo;
        """,
        params,
    )


def _top_table(source: list[dict[str, Any]], top_n: int | None) -> dict[str, Any]:
    """Build one independent Top-N table for a single greige fabric class.

    Page 4 now follows Page 5's split: flat woven and knit are ranked
    independently, so fabric type no longer needs to be repeated as a column.
    """
    total_qty = sum(_as_number(r.get("Quantity")) for r in source)
    total_amount = sum(_as_number(r.get("Amount")) for r in source)
    total_book = sum(_as_number(r.get("BookValue")) for r in source)
    total_provision = sum(_as_number(r.get("Provision")) for r in source)

    top = source if top_n is None else source[:top_n]
    output: list[dict[str, Any]] = []
    for index, row in enumerate(top, start=1):
        qty = _as_number(row.get("Quantity"))
        amount = _as_number(row.get("Amount"))
        book = _as_number(row.get("BookValue"))
        provision = _as_number(row.get("Provision"))
        output.append({
            "排行": index,
            "料件編號": row.get("ItemNo") or "",
            "庫齡": _format_age_span(row.get("MinAgeDays"), row.get("MaxAgeDays")),
            "單位": row.get("InventoryUnitCode") or "",
            "庫存數量": qty,
            "數量百分比": _pct(qty, total_qty),
            "存貨金額": amount,
            "金額百分比": _pct(amount, total_amount),
            "帳面價值": book,
            "提列率": _provision_rate(provision, amount),
        })

    top_qty = sum(_as_number(r.get("Quantity")) for r in top)
    top_amount = sum(_as_number(r.get("Amount")) for r in top)
    top_book = sum(_as_number(r.get("BookValue")) for r in top)
    top_provision = sum(_as_number(r.get("Provision")) for r in top)
    other_qty = total_qty - top_qty
    other_amount = total_amount - top_amount
    other_book = total_book - top_book
    other_provision = total_provision - top_provision

    if top_n is not None:
        output.extend([
            {
                "排行": f"前{top_n}大數量小計",
                "料件編號": "",
                "庫齡": "",
                "單位": "",
                "庫存數量": top_qty,
                "數量百分比": _pct(top_qty, total_qty),
                "存貨金額": top_amount,
                "金額百分比": _pct(top_amount, total_amount),
                "帳面價值": top_book,
                "提列率": _provision_rate(top_provision, top_amount),
            },
            {
                "排行": "其他",
                "料件編號": "",
                "庫齡": "",
                "單位": "",
                "庫存數量": other_qty,
                "數量百分比": _pct(other_qty, total_qty),
                "存貨金額": other_amount,
                "金額百分比": _pct(other_amount, total_amount),
                "帳面價值": other_book,
                "提列率": _provision_rate(other_provision, other_amount),
            },
        ])

    output.append({
        "排行": "總計",
        "料件編號": "",
        "庫齡": "",
        "單位": "",
        "庫存數量": total_qty,
        "數量百分比": 100.0 if total_qty else 0.0,
        "存貨金額": total_amount,
        "金額百分比": 100.0 if total_amount else 0.0,
        "帳面價值": total_book,
        "提列率": _provision_rate(total_provision, total_amount),
    })

    columns = [
        _column("排行"), _column("料件編號"), _column("庫齡"), _column("單位"),
        _column("庫存數量"), _column("數量百分比"),
        _column("存貨金額"), _column("金額百分比"),
        _column("帳面價值"), _column("提列率"),
    ]
    return {"columns": columns, "rows": output}


@router.post("/top-items")
def top_items(req: DashboardRequest) -> dict[str, Any]:
    """Page 4: greige Top-N, ranked separately for flat woven and knit.

    This intentionally uses the exact same ranking sources as Page 5 so both
    pages always select the same Top-N item set under the same filters.
    """
    period = _latest_period()
    top_n = None if req.top_n is None else max(1, min(int(req.top_n), 100))
    _, _, age_from, age_to, age_label = _top_age_filter(
        req.top_age_from, req.top_age_to, req.top_age_subrange
    )

    flat = _top_table(_top_source(req, period, "2"), top_n)
    knit = _top_table(_top_source(req, period, "1"), top_n)

    return {
        "report_period": period,
        "top_n": top_n,
        "top_n_label": "全部" if top_n is None else str(top_n),
        "type_code": "2",
        "type_label": "胚布",
        "age_from": age_from,
        "age_to": age_to,
        "age_label": age_label,
        "flat": flat,
        "knit": knit,
        # Keep the old key so older callers do not fail immediately.
        # New UI uses flat / knit and does not use items.
        "items": flat,
        "category_code": "2",
        "category_label": "胚布",
        "received_greige_exclusion_applied": False,
    }


def _responsibility_table(
    req: DashboardRequest,
    period: int,
    fabric_code: str,
    top_n: int | None,
) -> dict[str, Any]:
    # Ranking grain remains ItemNo within each greige fabric class, but the
    # exact same age From/To filter used on Page 4 is applied before ranking.
    ranked = _top_source(req, period, fabric_code)
    top = ranked if top_n is None else ranked[:top_n]
    top_items = [str(row.get("ItemNo") or "").strip() for row in top]
    top_items = [item for item in top_items if item]
    rank_map = {item: index for index, item in enumerate(top_items, start=1)}

    total_qty = sum(_as_number(r.get("Quantity")) for r in ranked)
    total_amount = sum(_as_number(r.get("Amount")) for r in ranked)

    columns = [
        _column("排行"), _column("料件編號"), _column("庫齡"), _column("單位"),
        _column("庫存數量"), _column("數量百分比"),
        _column("存貨金額"), _column("金額百分比"),
        _column("帳面價值"), _column("提列率"),
        _column("類別編號"), _column("品牌商"), _column("歸屬部門"),
    ]
    if not top_items:
        return {"columns": columns, "rows": []}

    where, params = _dashboard_where(
        req,
        periods=[period],
        force_types=["2"],
        force_fabric=fabric_code,
        three_years=False,
        apply_structure_filters=False,
    )
    age_sql, age_params, _, _, _ = _top_age_filter(req.top_age_from, req.top_age_to, req.top_age_subrange)
    where.append(age_sql)
    params.extend(age_params)

    placeholders = ",".join("?" for _ in top_items)
    where.append(f"CONVERT(nvarchar(200), ItemNo) IN ({placeholders})")
    params.extend(top_items)
    age_name, age_sort = _age_case_sql()

    source = _query(
        f"""
        SELECT
            CONVERT(nvarchar(200), ItemNo) AS ItemNo,
            CONVERT(nvarchar(40), InventoryUnitCode) AS InventoryUnitCode,
            {age_sort} AS AgeSort,
            {age_name} AS AgeBucket,
            MIN(AgeDays) AS MinAgeDays,
            MAX(AgeDays) AS MaxAgeDays,
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), SourceBatchNo))),N'') AS SourceBatchNo,
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), BrandOwnerName))),N'') AS BrandOwnerName,
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), AttributionDepartmentName))),N'') AS AttributionDepartmentName,
            SUM(ISNULL(InventoryQuantity,0)) AS Quantity,
            SUM(ISNULL(InventoryCostAmountNTD,0)) AS Amount,
            SUM(ISNULL(BookValueAmountNTD,0)) AS BookValue,
            SUM(ISNULL(ProvisionAmountNTD,0)) AS Provision
        FROM {FACT}
        WHERE {' AND '.join(where)}
        GROUP BY
            CONVERT(nvarchar(200), ItemNo),
            CONVERT(nvarchar(40), InventoryUnitCode),
            {age_sort}, {age_name},
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), SourceBatchNo))),N''),
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), BrandOwnerName))),N''),
            NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(200), AttributionDepartmentName))),N'')
        ORDER BY ItemNo, AgeSort, SourceBatchNo, BrandOwnerName, AttributionDepartmentName;
        """,
        params,
    )

    output: list[dict[str, Any]] = []
    for row in source:
        item = str(row.get("ItemNo") or "")
        qty = _as_number(row.get("Quantity"))
        amount = _as_number(row.get("Amount"))
        provision = _as_number(row.get("Provision"))
        output.append({
            "排行": rank_map.get(item, 999999),
            "料件編號": item,
            "庫齡": _format_age_span(row.get("MinAgeDays"), row.get("MaxAgeDays")),
            "單位": row.get("InventoryUnitCode") or "",
            "庫存數量": qty,
            "數量百分比": _pct(qty, total_qty),
            "存貨金額": amount,
            "金額百分比": _pct(amount, total_amount),
            "帳面價值": _as_number(row.get("BookValue")),
            "提列率": _provision_rate(provision, amount),
            "類別編號": row.get("SourceBatchNo") or "",
            "品牌商": row.get("BrandOwnerName") or "",
            "歸屬部門": row.get("AttributionDepartmentName") or "",
            "_AgeSort": int(row.get("AgeSort") or 99),
        })

    output.sort(
        key=lambda row: (
            int(row.get("排行") or 999999),
            int(row.get("_AgeSort") or 99),
            str(row.get("類別編號") or ""),
            str(row.get("品牌商") or ""),
            str(row.get("歸屬部門") or ""),
        )
    )
    for row in output:
        row.pop("_AgeSort", None)

    return {"columns": columns, "rows": output}


@router.post("/responsibility")
def responsibility(req: DashboardRequest) -> dict[str, Any]:
    period = _latest_period()
    top_n = None if req.top_n is None else max(1, min(int(req.top_n), 100))
    _, _, age_from, age_to, age_label = _top_age_filter(req.top_age_from, req.top_age_to, req.top_age_subrange)
    return {
        "report_period": period,
        "top_n": top_n,
        "top_n_label": "全部" if top_n is None else str(top_n),
        "type_code": "2",
        "type_label": "胚布",
        "age_from": age_from,
        "age_to": age_to,
        "age_label": age_label,
        "flat": _responsibility_table(req, period, "2", top_n),
        "knit": _responsibility_table(req, period, "1", top_n),
    }


# -----------------------------------------------------------------------------
# Dashboard Excel export (tabs 1-5)
# -----------------------------------------------------------------------------

def _excel_cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _write_dashboard_sheet(
    wb: Workbook,
    sheet_name: str,
    table: dict[str, Any],
    *,
    title: str | None = None,
    note: str | None = None,
) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])
    columns = table.get("columns") or []
    rows = table.get("rows") or []
    names = [str(item.get("name") if isinstance(item, dict) else item) for item in columns]
    labels = [
        str((item.get("label") or item.get("name")) if isinstance(item, dict) else item)
        for item in columns
    ]

    row_no = 1
    if title:
        ws.cell(row=row_no, column=1, value=title)
        ws.cell(row=row_no, column=1).font = Font(bold=True, size=14)
        if names:
            ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=len(names))
        row_no += 1
    if note:
        ws.cell(row=row_no, column=1, value=note)
        ws.cell(row=row_no, column=1).font = Font(italic=True)
        if names:
            ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=len(names))
        row_no += 1

    header_row = row_no
    for col_no, label in enumerate(labels, start=1):
        cell = ws.cell(row=header_row, column=col_no, value=label)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for source in rows:
        row_no += 1
        for col_no, name in enumerate(names, start=1):
            value = _excel_cell_value(source.get(name))
            cell = ws.cell(row=row_no, column=col_no, value=value)
            if isinstance(value, (int, float)):
                if "百分比" in name or "提列率" in name:
                    # API values are already percentages, e.g. 12.34 means 12.34%.
                    cell.number_format = '0.00"%"'
                else:
                    cell.number_format = '#,##0.00'

    if names:
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(names))}{max(header_row, row_no)}"
        for index, (name, label) in enumerate(zip(names, labels), start=1):
            sample = [len(str(label))]
            for source in rows[:300]:
                value = source.get(name)
                if value is not None:
                    sample.append(len(str(value)))
            ws.column_dimensions[get_column_letter(index)].width = min(max(sample) + 2, 32)


def _dashboard_excel_response(
    filename: str,
    sheets: list[tuple[str, dict[str, Any], str | None, str | None]],
) -> StreamingResponse:
    wb = Workbook()
    # Remove the empty default sheet after creating real sheets.
    default_ws = wb.active
    for sheet_name, table, title, note in sheets:
        _write_dashboard_sheet(wb, sheet_name, table, title=title, note=note)
    wb.remove(default_ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    encoded = quote(filename)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename=dashboard.xlsx; filename*=UTF-8''{encoded}"
            )
        },
    )


@router.post("/export/xlsx/{report_name}")
def export_dashboard_xlsx(report_name: str, req: DashboardRequest) -> StreamingResponse:
    report = str(report_name or "").strip().lower()
    now = datetime.now().strftime("%Y%m%d_%H%M%S")

    if report == "aging":
        data = aging(req)
        period = data.get("report_period")
        return _dashboard_excel_response(
            f"胚紗庫存分析_{period}_{now}.xlsx",
            [("胚紗庫存分析", data, "胚布＋紗庫存分析", f"ReportPeriod：{period}")],
        )

    if report == "category":
        data = category(req)
        period = data.get("report_period")
        return _dashboard_excel_response(
            f"庫齡結構摘要_{period}_{now}.xlsx",
            [("庫齡結構摘要", data, "紗／胚布庫齡結構摘要", f"ReportPeriod：{period}")],
        )

    if report == "comparison":
        data = period_comparison(req)
        if not data.get("available"):
            raise HTTPException(
                status_code=400,
                detail=data.get("message") or "目前無法執行兩期存貨差異。",
            )
        base_period = data.get("base_report_period")
        compare_period = data.get("compare_report_period")
        note = f"差異＝基準月份－比較月份；基準月份：{base_period}；比較月份：{compare_period}"
        return _dashboard_excel_response(
            f"兩期存貨差異_{base_period}_vs_{compare_period}_{now}.xlsx",
            [
                ("紗兩期差異", data["yarn"], "紗－兩期存貨差異", note),
                ("胚布兩期差異", data["greige"], "胚布－兩期存貨差異", note),
            ],
        )

    if report == "topn":
        data = top_items(req)
        period = data.get("report_period")
        top_n = data.get("top_n")
        top_label = data.get("top_n_label") or (str(top_n) if top_n is not None else "全部")
        age_label = data.get("age_label") or "0～2年"
        note = (
            f"ReportPeriod：{period}；類型：胚布；庫齡：{age_label}；Top {top_label}"
            "；平織／針織各自獨立排名"
        )
        export_scope = "全部" if top_n is None else f"前{top_n}大"
        return _dashboard_excel_response(
            f"胚布{export_scope}_平織針織_{age_label}_{period}_{now}.xlsx",
            [
                ("平織TopN", data["flat"], f"胚布{export_scope}－平織－{age_label}", note),
                ("針織TopN", data["knit"], f"胚布{export_scope}－針織－{age_label}", note),
            ],
        )

    if report == "responsibility":
        data = responsibility(req)
        period = data.get("report_period")
        top_n = data.get("top_n")
        top_label = data.get("top_n_label") or (str(top_n) if top_n is not None else "全部")
        age_label = data.get("age_label") or "0～2年"
        note = f"ReportPeriod：{period}；類型：胚布；庫齡：{age_label}；Top {top_label}"
        export_scope = "全部" if top_n is None else f"前{top_n}大"
        return _dashboard_excel_response(
            f"胚布{export_scope}責任歸屬_{age_label}_{period}_{now}.xlsx",
            [
                ("平織責任歸屬", data["flat"], f"胚布{export_scope}－平織責任歸屬－{age_label}", note),
                ("針織責任歸屬", data["knit"], f"胚布{export_scope}－針織責任歸屬－{age_label}", note),
            ],
        )

    raise HTTPException(status_code=400, detail="此頁籤不支援 Dashboard Excel 匯出。")
